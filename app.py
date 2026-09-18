import os
from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import text, or_
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from email.mime.image import MIMEImage
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta
import logging
import io
import re
import zipfile
import urllib.request
import urllib.parse
import http.cookiejar
import html
import xml.etree.ElementTree as ET
from werkzeug.utils import secure_filename

# ===================================
# Initialize Flask App
# ===================================
app = Flask(__name__, static_folder='static')
CORS(app, resources={
    r"/api/*": {"origins": "*"},
    r"/static/*": {"origins": "*"}
})

# ===================================
# Database Configuration
# ===================================
db_url = os.environ.get('DATABASE_URL')
if not db_url:
    raise RuntimeError("DATABASE_URL environment variable is not set")

# Adjust for postgres:// vs postgresql://
if db_url.startswith('postgres://'):
    db_url = db_url.replace('postgres://', 'postgresql://', 1)

app.config.update({
    'SQLALCHEMY_DATABASE_URI': db_url,
    'SQLALCHEMY_TRACK_MODIFICATIONS': False,
    'SQLALCHEMY_ENGINE_OPTIONS': {
        'pool_size': 5,
        'pool_recycle': 300,
        'pool_pre_ping': True
    },
    'UPLOAD_FOLDER': os.path.join(os.getcwd(), 'uploads'),
    'MAX_CONTENT_LENGTH': 16 * 1024 * 1024  # 16MB
})

os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
db = SQLAlchemy(app)

# ===================================
# Database Models
# ===================================
class RepairReport(db.Model):
    __tablename__ = 'repair_reports'
    id = db.Column(db.Integer, primary_key=True)
    container_number = db.Column(db.String(11), nullable=False)
    report_date = db.Column(db.Date, nullable=False)
    technician_name = db.Column(db.String(100), nullable=False)
    model = db.Column(db.String(100))
    model_family = db.Column(db.String(50))
    serial_number = db.Column(db.String(100))
    warranty_id = db.Column(db.String(100))
    warranty_status = db.Column(db.String(100))
    setpoint = db.Column(db.Float)
    vents = db.Column(db.String(50))
    humidity = db.Column(db.String(50))
    ambient_temp = db.Column(db.Float)
    supply_temp_before = db.Column(db.Float)
    supply_temp_after = db.Column(db.Float)
    return_temp_before = db.Column(db.Float)
    return_temp_after = db.Column(db.Float)
    temp_in_range = db.Column(db.String(50))
    problem_description = db.Column(db.Text)
    comments = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

class RepairJob(db.Model):
    __tablename__ = 'repair_jobs'
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey('repair_reports.id'), nullable=False)
    job_code = db.Column(db.String(50))
    description = db.Column(db.String(255))
    part_number = db.Column(db.String(100))
    part_description = db.Column(db.String(255))
    quantity = db.Column(db.Float)
    damage_type = db.Column(db.String(50))
    old_serial = db.Column(db.String(100))
    new_serial = db.Column(db.String(100))
    bottle_number = db.Column(db.String(100))
    labor_hours = db.Column(db.Float)
    freon_sent_at = db.Column(db.DateTime)

class JobCodeMaster(db.Model):
    __tablename__ = 'job_code_master'
    id = db.Column(db.Integer, primary_key=True)
    job_code = db.Column(db.String(50), unique=True, nullable=False, index=True)
    description = db.Column(db.String(255), nullable=False, index=True)
    job_code_type = db.Column(db.String(50), nullable=False)
    component_code = db.Column(db.String(50))
    repair_type = db.Column(db.String(50))
    damage_location = db.Column(db.String(50))
    damage_type = db.Column(db.String(50))
    material_type = db.Column(db.String(50))

class SparePartMaster(db.Model):
    __tablename__ = 'spare_part_master'
    id = db.Column(db.Integer, primary_key=True)
    part_number = db.Column(db.String(100), unique=True, nullable=False, index=True)
    replaced_by = db.Column(db.String(100))
    description = db.Column(db.String(500), index=True)
    max_qty = db.Column(db.String(50))

class LaborRuleMaster(db.Model):
    __tablename__ = 'labor_rule_master'
    id = db.Column(db.Integer, primary_key=True)
    model_family = db.Column(db.String(50), nullable=False, index=True)
    part_number = db.Column(db.String(100), nullable=False, index=True)
    part_description = db.Column(db.Text)
    damage_code = db.Column(db.String(50), index=True)
    repair_code = db.Column(db.String(50), index=True)
    repair_description = db.Column(db.Text)
    first_hour = db.Column(db.Float, nullable=False, default=0)

class Alarm(db.Model):
    __tablename__ = 'alarms'
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey('repair_reports.id'), nullable=False)
    alarm_code = db.Column(db.String(100))

class RepairListItem(db.Model):
    __tablename__ = 'repair_list_items'
    id = db.Column(db.Integer, primary_key=True)
    container_number = db.Column(db.String(11), unique=True, nullable=False, index=True)
    order_temp = db.Column(db.String(50))
    position = db.Column(db.String(100))
    alarms = db.Column(db.Text)
    etd = db.Column(db.String(50))
    vessel = db.Column(db.String(150))
    requested_repair = db.Column(db.Text)
    remarks = db.Column(db.Text)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)

    def to_dict(self):
        return {
            "id": self.id,
            "container_number": self.container_number,
            "order_temp": self.order_temp or "",
            "position": self.position or "",
            "alarms": self.alarms or "",
            "etd": self.etd or "",
            "vessel": self.vessel or "",
            "requested_repair": self.requested_repair or "",
            "remarks": self.remarks or ""
        }

class RepairZoneRequest(db.Model):
    __tablename__ = 'repair_zone_requests'
    id = db.Column(db.Integer, primary_key=True)
    container_number = db.Column(db.String(11), nullable=False, index=True)
    technician_name = db.Column(db.String(100), nullable=False)
    current_position = db.Column(db.String(100), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

def read_local_excel_sheet(file_path, sheet_name):
    """Read a worksheet from a bundled .xlsx file without an external Excel dependency."""
    spreadsheet_ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    relationship_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    office_rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

    def local_excel_column_index(cell_reference):
        letters = re.match(r'[A-Z]+', cell_reference or '')
        if not letters:
            return 0
        result = 0
        for char in letters.group(0):
            result = result * 26 + ord(char) - 64
        return result - 1

    with zipfile.ZipFile(file_path) as archive:
        shared_strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            shared_root = ET.fromstring(archive.read('xl/sharedStrings.xml'))
            for item in shared_root.findall(f'{{{spreadsheet_ns}}}si'):
                shared_strings.append(''.join(node.text or '' for node in item.iter(f'{{{spreadsheet_ns}}}t')))

        workbook_root = ET.fromstring(archive.read('xl/workbook.xml'))
        selected_sheet = None
        for sheet in workbook_root.findall(f'.//{{{spreadsheet_ns}}}sheet'):
            if sheet.attrib.get('name', '').strip().casefold() == sheet_name.strip().casefold():
                selected_sheet = sheet
                break
        if selected_sheet is None:
            raise ValueError(f'Worksheet {sheet_name!r} not found in {os.path.basename(file_path)}')

        relationship_id = selected_sheet.attrib.get(f'{{{office_rel_ns}}}id')
        rels_root = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        target = None
        for relation in rels_root.findall(f'{{{relationship_ns}}}Relationship'):
            if relation.attrib.get('Id') == relationship_id:
                target = relation.attrib.get('Target')
                break
        if not target:
            raise ValueError(f'Could not locate worksheet {sheet_name!r}')

        worksheet_path = target.lstrip('/') if target.startswith('/xl/') else 'xl/' + target.lstrip('/')
        worksheet_path = worksheet_path.replace('xl/xl/', 'xl/')
        sheet_root = ET.fromstring(archive.read(worksheet_path))
        rows = []
        for row in sheet_root.findall(f'.//{{{spreadsheet_ns}}}row'):
            values = {}
            next_column = 0
            for cell in row.findall(f'{{{spreadsheet_ns}}}c'):
                cell_reference = cell.attrib.get('r', '')
                column = local_excel_column_index(cell_reference) if cell_reference else next_column
                next_column = column + 1
                cell_type = cell.attrib.get('t')
                value_node = cell.find(f'{{{spreadsheet_ns}}}v')
                if cell_type == 'inlineStr':
                    inline = cell.find(f'{{{spreadsheet_ns}}}is')
                    value = ''.join(node.text or '' for node in inline.iter(f'{{{spreadsheet_ns}}}t')) if inline is not None else ''
                else:
                    value = value_node.text if value_node is not None else ''
                    if cell_type == 's' and value:
                        value = shared_strings[int(value)]
                values[column] = value
            if values:
                rows.append([values.get(i, '') for i in range(max(values) + 1)])
    return rows


def seed_master_data():
    """Synchronize approved job codes and load model-specific parts/labor masters."""
    base_dir = os.path.dirname(os.path.abspath(__file__))
    job_file = os.path.join(base_dir, 'job_codes_master.xlsx')
    parts_file = os.path.join(base_dir, 'parts_master.xlsx')

    # Synchronize the database with the compact job-code master.
    # This replaces legacy REMS codes that are no longer in the approved list.
    if not os.path.exists(job_file):
        app.logger.warning('Job-code master file not found: %s', job_file)
    else:
        rows = read_local_excel_sheet(job_file, 'Job Code Master')
        if rows:
            headers = {str(v).strip().casefold(): i for i, v in enumerate(rows[0])}

            def val(row, name):
                i = headers.get(name.casefold())
                return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ''

            source_codes = {}
            for row in rows[1:]:
                code = val(row, 'Job Code').upper()
                if not code:
                    continue
                source_codes[code] = {
                    'description': val(row, 'Job Code Description'),
                    'job_code_type': val(row, 'Job Code Type') or 'Terminal ER',
                    'component_code': val(row, 'Component Code'),
                    'repair_type': val(row, 'Repair Type'),
                    'material_type': val(row, 'Material Type')
                }

            existing_codes = {item.job_code: item for item in JobCodeMaster.query.all()}

            for code, data in source_codes.items():
                item = existing_codes.get(code)
                if item is None:
                    item = JobCodeMaster(job_code=code)
                    db.session.add(item)
                item.description = data['description']
                item.job_code_type = data['job_code_type']
                item.component_code = data['component_code']
                item.repair_type = data['repair_type']
                item.material_type = data['material_type']
                item.damage_location = ''
                item.damage_type = ''

            for code, item in existing_codes.items():
                if code not in source_codes:
                    db.session.delete(item)

            db.session.commit()
            app.logger.info('Synchronized %s approved job codes', len(source_codes))

    model_files = {
        'PrimeLine': 'Primeline.xlsx',
        'OptimaLine': 'Optimaline.xlsx',
        'StarCool': 'Starcool.xlsx',
        'ThinLINE': 'Thinline.xlsx',
        'EliteLINE': 'Eliteline.xlsx',
        'NaturaLINE': 'Naturaline.xlsx',
    }

    # The six model-specific masters are the source for both parts and First Hour labor.
    # No OneDrive or monthly parts workbook is required.
    if SparePartMaster.query.count() == 0 or LaborRuleMaster.query.count() == 0:
        existing_parts = {p.part_number for p in SparePartMaster.query.all()}
        existing_rules = {(r.model_family, r.part_number, r.damage_code, r.repair_code)
                          for r in LaborRuleMaster.query.all()}
        parts_added = 0
        rules_added = 0
        for model_family, filename in model_files.items():
            file_path = os.path.join(base_dir, filename)
            if not os.path.exists(file_path):
                app.logger.warning('Model labor master not found: %s', file_path)
                continue
            rows = read_local_excel_sheet(file_path, 'Export')
            if not rows:
                continue
            headers = {str(v).strip().casefold(): i for i, v in enumerate(rows[0])}
            def val(row, name):
                i = headers.get(name.casefold())
                return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ''
            for row in rows[1:]:
                part_number = val(row, 'part_code').upper()
                description = val(row, 'Part_Description')
                damage_code = val(row, 'DamageCode').upper()
                repair_code = val(row, 'RepairCode').upper()
                repair_description = val(row, 'Repair_description')

                if not part_number:
                    continue

                # Reject malformed workbook rows instead of letting corrupted cell
                # contents break PostgreSQL during application startup.
                if len(part_number) > 100:
                    app.logger.warning(
                        "Skipping malformed part number in %s: length=%s, starts_with=%r",
                        filename, len(part_number), part_number[:80]
                    )
                    continue
                if len(damage_code) > 50 or len(repair_code) > 50:
                    app.logger.warning(
                        "Skipping malformed labor rule in %s for part %s",
                        filename, part_number
                    )
                    continue

                description = description[:5000]
                repair_description = repair_description[:5000]
                if part_number not in existing_parts:
                    db.session.add(SparePartMaster(
                        part_number=part_number, replaced_by='', description=description, max_qty=''
                    ))
                    existing_parts.add(part_number)
                    parts_added += 1
                first_hour_text = val(row, 'FirstHour')
                try:
                    first_hour = float(first_hour_text or 0)
                except (TypeError, ValueError):
                    first_hour = 0.0
                rule_key = (model_family, part_number, damage_code, repair_code)
                if rule_key not in existing_rules:
                    db.session.add(LaborRuleMaster(
                        model_family=model_family,
                        part_number=part_number,
                        part_description=description,
                        damage_code=damage_code,
                        repair_code=repair_code,
                        repair_description=repair_description,
                        first_hour=first_hour
                    ))
                    existing_rules.add(rule_key)
                    rules_added += 1
            db.session.flush()
        db.session.commit()
        app.logger.info('Loaded %s model spare parts and %s labor rules', parts_added, rules_added)

# Initialize DB
with app.app_context():
    try:
        db.create_all()
        # Existing REMS databases may still have repair_jobs.quantity as INTEGER.
        # PostgreSQL safely converts existing integer quantities to double precision.
        try:
            db.session.execute(text("ALTER TABLE repair_jobs ALTER COLUMN quantity TYPE DOUBLE PRECISION USING quantity::double precision"))
            db.session.commit()
        except Exception:
            db.session.rollback()

        # Add fields introduced by the REMS One Vision workflow to existing databases.
        try:
            db.session.execute(text(
                "ALTER TABLE repair_reports ADD COLUMN IF NOT EXISTS model_family VARCHAR(50)"
            ))
            db.session.execute(text(
                "ALTER TABLE repair_jobs ADD COLUMN IF NOT EXISTS bottle_number VARCHAR(100)"
            ))
            db.session.execute(text(
                "ALTER TABLE repair_jobs ADD COLUMN IF NOT EXISTS freon_sent_at TIMESTAMP"
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()

        # Upgrade description columns on an existing REMS database before seeding.
        try:
            db.session.execute(text(
                "ALTER TABLE labor_rule_master "
                "ALTER COLUMN part_description TYPE TEXT, "
                "ALTER COLUMN repair_description TYPE TEXT"
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()

        seed_master_data()
        db.session.execute(text("SELECT 1"))
        app.logger.info("Database initialized successfully")
    except Exception as e:
        app.logger.critical(f"Database initialization failed: {str(e)}")
        raise

# ===================================
# Routes
# ===================================

@app.route('/')
def serve_index():
    return send_from_directory('.', 'index.html')

@app.route('/favicon.ico')
def favicon():
    return send_from_directory('static', 'favicon.ico', mimetype='image/vnd.microsoft.icon')

@app.route('/static/<path:path>')
def serve_static(path):
    return send_from_directory('static', path)

# --------- LOGIN ROUTE ----------
@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json()
    username = data.get("username")
    password = data.get("password")

    # Example ENV: TECHNICIANS="Admin:admin123,Brahim:bm123"
    valid_users = {}
    env_users = os.environ.get("TECHNICIANS", "")
    for pair in env_users.split(","):
        if ":" in pair:
            user, pwd = pair.split(":", 1)
            valid_users[user.strip()] = pwd.strip()

    if username in valid_users and valid_users[username] == password:
        return jsonify({"status": "success"})
    return jsonify({"status": "error", "message": "Invalid credentials"}), 401

# --------- MASTER DATA ----------
@app.route('/api/job-codes', methods=['GET'])
def search_job_codes():
    q = str(request.args.get('q') or '').strip()
    query = JobCodeMaster.query
    if q:
        pattern = f"%{q}%"
        query = query.filter(or_(
            JobCodeMaster.job_code.ilike(pattern),
            JobCodeMaster.description.ilike(pattern)
        ))
    jobs = query.order_by(JobCodeMaster.job_code.asc()).limit(50).all()
    results = [{
        "id": job.job_code,
        "text": f"{job.job_code} — {job.description}",
        "code": job.job_code,
        "description": job.description,
        "type": job.job_code_type,
        "component_code": job.component_code or "",
        "repair_type": job.repair_type or "",
        "material_type": job.material_type or ""
    } for job in jobs]

    # REMS-only special job; no master-data/database change required.
    if not q or "ice" in q.casefold() or "procedure" in q.casefold():
        results.insert(0, {
            "id": "ICE_PROCEDURE",
            "text": "Ice Procedure",
            "code": "ICE_PROCEDURE",
            "description": "Ice Procedure",
            "type": "special",
            "component_code": "",
            "repair_type": "",
            "material_type": ""
        })

    return jsonify({"results": results[:50]})

@app.route('/api/parts', methods=['GET'])
def search_spare_parts():
    q = str(request.args.get('q') or '').strip()
    model_family = str(request.args.get('model_family') or '').strip()

    # The model-specific Excel files are loaded into LaborRuleMaster.
    # Query that table directly so only parts belonging to the selected
    # Model Family can ever be returned.
    if not model_family:
        return jsonify({"results": []})

    query = LaborRuleMaster.query.filter(
        LaborRuleMaster.model_family == model_family
    )

    if q:
        pattern = f"%{q}%"
        query = query.filter(or_(
            LaborRuleMaster.part_number.ilike(pattern),
            LaborRuleMaster.part_description.ilike(pattern)
        ))

    rules = (
        query
        .order_by(LaborRuleMaster.part_number.asc())
        .limit(250)
        .all()
    )

    # A part can have several damage/repair rules in the same model file.
    # Return each part number only once to the technician.
    results = []
    seen = set()
    for rule in rules:
        part_number = (rule.part_number or '').strip()
        if not part_number or part_number in seen:
            continue
        seen.add(part_number)
        description = (rule.part_description or '').strip()
        results.append({
            "id": part_number,
            "text": f"{part_number} — {description}",
            "part_number": part_number,
            "description": description,
            "replaced_by": "",
            "max_qty": ""
        })
        if len(results) >= 50:
            break

    return jsonify({
        "results": results,
        "model_family": model_family
    })

@app.route('/api/labor-hours', methods=['GET'])
def get_labor_hours():
    model_text = str(request.args.get('model') or '').strip().casefold()
    part_number = str(request.args.get('part_number') or '').strip().upper()
    repair_code = str(request.args.get('repair_code') or '').strip().upper()
    quantity_text = str(request.args.get('quantity') or '1').strip()
    try:
        quantity = float(quantity_text or 1)
    except ValueError:
        quantity = 1.0

    aliases = [
        ('optimaline', 'OptimaLine'), ('optima line', 'OptimaLine'),
        ('primeline', 'PrimeLine'), ('prime line', 'PrimeLine'),
        ('starcool', 'StarCool'), ('star cool', 'StarCool'),
        ('thinline', 'ThinLINE'), ('thin line', 'ThinLINE'),
        ('eliteline', 'EliteLINE'), ('elite line', 'EliteLINE'),
        ('naturaline', 'NaturaLINE'), ('natura line', 'NaturaLINE'),
    ]
    model_family = next((family for alias, family in aliases if alias in model_text), '')
    if not model_family or not part_number:
        return jsonify({'status': 'not_found', 'first_hour': 0, 'labor_hours': 0})

    query = LaborRuleMaster.query.filter_by(model_family=model_family, part_number=part_number)
    # Current REMS/One Vision rule: Damage Type is always Broken (BR).
    broken = query.filter(LaborRuleMaster.damage_code == 'BR')
    if repair_code:
        exact = broken.filter(LaborRuleMaster.repair_code == repair_code).first()
        rule = exact or broken.first()
    else:
        rule = broken.first()
    if rule is None:
        rule = query.first()
    if rule is None:
        return jsonify({'status': 'not_found', 'model_family': model_family, 'first_hour': 0, 'labor_hours': 0})

    labor_hours = round((rule.first_hour or 0) * quantity, 2)
    is_cable = 'cable' in (rule.part_description or '').casefold()
    return jsonify({
        'status': 'success',
        'model_family': model_family,
        'first_hour': rule.first_hour or 0,
        'labor_hours': labor_hours,
        'repair_code': rule.repair_code or '',
        'is_cable': is_cable,
        'quantity_note': '1 roll = 18 m' if is_cable else ''
    })

# --------- REPAIR LIST ----------
@app.route('/api/repair-list', methods=['GET'])
def get_repair_list():
    excel_share_url = os.environ.get('EXCEL_SHARE_URL', '').strip()
    if excel_share_url:
        try:
            items = read_repair_list_from_excel(excel_share_url)
            add_recent_history_counts(items)
            return jsonify({"status": "success", "items": items, "source": "excel"})
        except Exception as exc:
            app.logger.error("Excel repair-list sync failed: %s", exc, exc_info=True)

    # Keep REMS usable if OneDrive is temporarily unavailable or not configured.
    items = RepairListItem.query.order_by(RepairListItem.created_at.asc()).all()
    item_data = [item.to_dict() for item in items]
    add_recent_history_counts(item_data)
    return jsonify({"status": "success", "items": item_data, "source": "database"})

@app.route('/api/report-history/<string:container_nr>', methods=['GET'])
def get_report_history(container_nr):
    container_nr = container_nr.strip().upper()
    cutoff_date = (datetime.utcnow() - timedelta(days=30)).date()
    reports = RepairReport.query.filter(
        RepairReport.container_number == container_nr,
        RepairReport.report_date >= cutoff_date
    ).order_by(RepairReport.report_date.desc(), RepairReport.id.desc()).all()

    history = []
    for report in reports:
        jobs = RepairJob.query.filter_by(report_id=report.id).order_by(RepairJob.id.asc()).all()
        alarms = Alarm.query.filter_by(report_id=report.id).order_by(Alarm.id.asc()).all()
        history.append({
            "id": report.id,
            "report_date": report.report_date.strftime('%d/%m/%Y'),
            "technician_name": report.technician_name or "",
            "problem_description": report.problem_description or "",
            "comments": report.comments or "",
            "alarms": [alarm.alarm_code for alarm in alarms if alarm.alarm_code],
            "jobs": [{
                "job_code": job.job_code or "",
                "description": job.description or "",
                "part_number": job.part_number or "",
                "part_description": job.part_description or "",
                "quantity": job.quantity or 0
            } for job in jobs]
        })
    return jsonify({"status": "success", "container_number": container_nr, "history": history})

@app.route('/api/repair-list', methods=['POST'])
def upsert_repair_list_item():
    """Create/update a work item. The existing email automation can call this endpoint."""
    data = request.get_json(silent=True) or request.form
    container_nr = str(data.get('container_number') or data.get('containernr') or '').strip().upper()
    if not (len(container_nr) == 11 and container_nr[:4].isalpha() and container_nr[4:].isdigit()):
        return jsonify({"status": "error", "message": "Invalid container number format"}), 400

    item = RepairListItem.query.filter_by(container_number=container_nr).first()
    if item is None:
        item = RepairListItem(container_number=container_nr)
        db.session.add(item)

    for field in ('order_temp', 'position', 'alarms', 'etd', 'vessel', 'requested_repair', 'remarks'):
        if field in data:
            setattr(item, field, str(data.get(field) or '').strip())

    db.session.commit()
    return jsonify({"status": "success", "item": item.to_dict()}), 200

@app.route('/api/repair-list/<string:container_nr>', methods=['DELETE'])
def delete_repair_list_item(container_nr):
    item = RepairListItem.query.filter_by(container_number=container_nr.strip().upper()).first()
    if item is None:
        return jsonify({"status": "error", "message": "Container not found"}), 404
    db.session.delete(item)
    db.session.commit()
    return jsonify({"status": "success"})

@app.route('/api/request-repair-zone', methods=['POST'])
def request_repair_zone():
    data = request.get_json(silent=True) or {}
    container_nr = str(data.get('container_number') or '').strip().upper()
    technician_name = str(data.get('technician_name') or '').strip()
    current_position = str(data.get('current_position') or '').strip()

    if not (len(container_nr) == 11 and container_nr[:4].isalpha() and container_nr[4:].isdigit()):
        return jsonify({"status": "error", "message": "Invalid container number format"}), 400
    if not technician_name:
        return jsonify({"status": "error", "message": "Technician name is required"}), 400
    if not current_position:
        return jsonify({"status": "error", "message": "Current position is required"}), 400

    duplicate_since = datetime.utcnow() - timedelta(minutes=15)
    duplicate = RepairZoneRequest.query.filter(
        RepairZoneRequest.container_number == container_nr,
        RepairZoneRequest.created_at >= duplicate_since
    ).first()
    if duplicate:
        return jsonify({
            "status": "error",
            "message": "A repair-zone request for this reefer was already sent in the last 15 minutes"
        }), 409

    request_record = RepairZoneRequest(
        container_number=container_nr,
        technician_name=technician_name,
        current_position=current_position
    )
    db.session.add(request_record)

    try:
        send_repair_zone_email(container_nr, technician_name, current_position)
        db.session.commit()
    except Exception as exc:
        db.session.rollback()
        app.logger.error("Repair-zone email failed: %s", exc, exc_info=True)
        return jsonify({"status": "error", "message": "The request email could not be sent"}), 502

    return jsonify({
        "status": "success",
        "message": f"Repair-zone request sent for {container_nr}"
    })

# --------- SUBMIT REPORT ----------
@app.route('/api/submit', methods=['POST'])
def submit_report():
    try:
        if not request.is_json and not request.form:
            return jsonify({"status": "error", "message": "Unsupported content type"}), 415

        form_data = request.form if request.form else request.get_json()
        files = request.files

        # Validate container number
        container_nr = form_data.get('containernr', '')
        if not (len(container_nr) == 11 and container_nr[:4].isalpha() and container_nr[4:].isdigit()):
            return jsonify({"status": "error", "message": "Invalid container number format"}), 400

        # Create Repair Report
        report = RepairReport(
            container_number=container_nr,
            report_date=datetime.strptime(form_data.get('datum'), '%Y-%m-%d').date(),
            technician_name=form_data.get('naam'),
            model=form_data.get('model'),
            model_family=form_data.get('model_family'),
            serial_number=form_data.get('serienr'),
            warranty_id=form_data.get('warranty_id'),
            warranty_status=form_data.get('garantie'),
            setpoint=float(form_data.get('setpoint', 0)),
            vents=form_data.get('vents'),
            humidity=form_data.get('hum'),
            ambient_temp=float(form_data.get('ambient', 0)),
            supply_temp_before=float(form_data.get('supply_voor', 0)),
            supply_temp_after=float(form_data.get('supply_na', 0)),
            return_temp_before=float(form_data.get('return_voor', 0)),
            return_temp_after=float(form_data.get('return_na', 0)),
            temp_in_range=form_data.get('temp_in_range'),
            problem_description=form_data.get('probleem'),
            comments=form_data.get('opmerkingen')
        )
        db.session.add(report)
        db.session.flush()  # Get report ID

        # Jobs
        # Ice Procedure is email-only and is not stored in repair_jobs.
        job_count = int(form_data.get('job_count', 0))
        ice_procedures = []
        for i in range(job_count):
            job_code = str(form_data.get(f'job[{i}][code]') or '').strip().upper()

            if job_code == 'ICE_PROCEDURE':
                alarm_active = str(form_data.get(f'job[{i}][ice_alarm_active]') or '').strip()
                ice_alarm = str(form_data.get(f'job[{i}][ice_alarm]') or '').strip()
                ice_data = {
                    'cd26': str(form_data.get(f'job[{i}][ice_cd26]') or '').strip(),
                    'cd27': str(form_data.get(f'job[{i}][ice_cd27]') or '').strip(),
                    'drain_hose': str(form_data.get(f'job[{i}][ice_drain_hose]') or '').strip(),
                    'upper_supply_sensor': str(form_data.get(f'job[{i}][ice_upper_supply_sensor]') or '').strip(),
                    'lower_supply_sensor': str(form_data.get(f'job[{i}][ice_lower_supply_sensor]') or '').strip(),
                    'alarm_active': alarm_active,
                    'alarm': ice_alarm
                }
                if not all([
                    ice_data['cd26'], ice_data['cd27'], ice_data['drain_hose'],
                    ice_data['upper_supply_sensor'], ice_data['lower_supply_sensor'],
                    ice_data['alarm_active']
                ]):
                    return jsonify({"status": "error", "message": "Please complete all Ice Procedure fields"}), 400
                if alarm_active.casefold() == 'yes' and not ice_alarm:
                    return jsonify({"status": "error", "message": "Please enter the active alarm for the Ice Procedure"}), 400
                ice_procedures.append(ice_data)
                continue

            job = RepairJob(
                report_id=report.id,
                job_code=form_data.get(f'job[{i}][code]'),
                description=form_data.get(f'job[{i}][description]'),
                part_number=form_data.get(f'job[{i}][part_number]'),
                part_description=form_data.get(f'job[{i}][part_description]'),
                quantity=float(form_data.get(f'job[{i}][quantity]') or 1),
                damage_type=form_data.get(f'job[{i}][damage_type]'),
                old_serial=form_data.get(f'job[{i}][old_serial]'),
                new_serial=form_data.get(f'job[{i}][new_serial]'),
                bottle_number=form_data.get(f'job[{i}][bottle_number]'),
                labor_hours=float(form_data.get(f'job[{i}][labor_hours]') or 0)
            )
            db.session.add(job)

        # Alarms
        for alarm in request.form.getlist('alarm[]'):
            if alarm.strip():
                db.session.add(Alarm(report_id=report.id, alarm_code=alarm.strip()))

        # Files — save every selected photo and keep Before/After counts for the email.
        saved_files = []
        photo_counts = {"before": 0, "after": 0}
        file_counter = 0
        for file_key in files.keys():
            for file in files.getlist(file_key):
                if file and allowed_file(file.filename):
                    original_name = secure_filename(file.filename)
                    name, extension = os.path.splitext(original_name)
                    filename = f"report_{report.id}_{file_counter}_{name}{extension}"
                    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                    file.save(filepath)
                    saved_files.append(filepath)

                    key_lower = str(file_key).lower()
                    if "fotos_voor" in key_lower:
                        photo_counts["before"] += 1
                    elif "fotos_na" in key_lower:
                        photo_counts["after"] += 1

                    file_counter += 1

        # Send Email
        try:
            # Get jobs and alarms for this report
            jobs = RepairJob.query.filter_by(report_id=report.id).all()
            alarms = Alarm.query.filter_by(report_id=report.id).all()
            
            send_email(
                subject=container_nr,
                body="Repair Report submitted", 
                attachments=saved_files,
                report=report,
                jobs=jobs,
                alarms=alarms,
                afmelding=form_data.get("afmelding", ""),
                photo_counts=photo_counts,
                ice_procedures=ice_procedures
            )
        except Exception as e:
            app.logger.error(f"Email failed: {str(e)}")

        # A completed repair disappears from the technician work list.
        # Incomplete reports (Afmelding = Nee) remain available for follow-up.
        if str(form_data.get("afmelding", "")).strip().lower() == "ja":
            RepairListItem.query.filter_by(container_number=container_nr.upper()).delete()

        db.session.commit()
        return jsonify({"status": "success", "message": "Report submitted successfully", "report_id": report.id})

    except Exception as e:
        db.session.rollback()
        app.logger.error(f"Submission failed: {str(e)}", exc_info=True)
        return jsonify({"status": "error", "message": "Internal server error"}), 500

# ===================================
# Helpers
# ===================================
def add_recent_history_counts(items):
    container_numbers = {str(item.get('container_number', '')).strip().upper() for item in items}
    container_numbers.discard('')
    if not container_numbers:
        return
    cutoff_date = (datetime.utcnow() - timedelta(days=30)).date()
    reports = RepairReport.query.with_entities(RepairReport.container_number).filter(
        RepairReport.container_number.in_(container_numbers),
        RepairReport.report_date >= cutoff_date
    ).all()
    counts = {}
    for (container_number,) in reports:
        key = (container_number or '').upper()
        counts[key] = counts.get(key, 0) + 1
    for item in items:
        item['history_count'] = counts.get(str(item.get('container_number', '')).upper(), 0)

def excel_column_index(cell_reference):
    letters = re.match(r'[A-Z]+', cell_reference or '')
    if not letters:
        return 0
    result = 0
    for char in letters.group(0):
        result = result * 26 + ord(char) - 64
    return result - 1

def excel_date_display(value):
    try:
        serial = float(value)
        if serial <= 0:
            return str(value or '')
        date_value = datetime(1899, 12, 30) + __import__('datetime').timedelta(days=serial)
        return date_value.strftime('%d/%m')
    except (TypeError, ValueError, OverflowError):
        return str(value or '').strip()

def read_repair_list_from_excel(share_url):
    """Download a view-only OneDrive workbook and read its first worksheet."""
    parsed = urllib.parse.urlsplit(share_url)
    query = dict(urllib.parse.parse_qsl(parsed.query, keep_blank_values=True))
    query['download'] = '1'
    download_url = urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path, urllib.parse.urlencode(query), parsed.fragment))
    request_obj = urllib.request.Request(download_url, headers={
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/122 Safari/537.36',
        'Accept': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream;q=0.9,*/*;q=0.8'
    })
    cookie_jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cookie_jar))
    with opener.open(request_obj, timeout=20) as response:
        workbook_bytes = response.read(20 * 1024 * 1024 + 1)
    if len(workbook_bytes) > 20 * 1024 * 1024:
        raise ValueError('Excel file exceeds the 20 MB safety limit')
    if not workbook_bytes.startswith(b'PK'):
        raise ValueError('The sharing link did not return an Excel workbook')

    spreadsheet_ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    relationship_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    office_rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

    with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as archive:
        shared_strings = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            shared_root = ET.fromstring(archive.read('xl/sharedStrings.xml'))
            for item in shared_root.findall(f'{{{spreadsheet_ns}}}si'):
                shared_strings.append(''.join(node.text or '' for node in item.iter(f'{{{spreadsheet_ns}}}t')))

        workbook_root = ET.fromstring(archive.read('xl/workbook.xml'))
        first_sheet = workbook_root.find(f'.//{{{spreadsheet_ns}}}sheet')
        if first_sheet is None:
            return []
        relationship_id = first_sheet.attrib.get(f'{{{office_rel_ns}}}id')
        rels_root = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        target = None
        for relation in rels_root.findall(f'{{{relationship_ns}}}Relationship'):
            if relation.attrib.get('Id') == relationship_id:
                target = relation.attrib.get('Target')
                break
        if not target:
            raise ValueError('Could not locate the Excel worksheet')
        worksheet_path = target.lstrip('/') if target.startswith('/xl/') else 'xl/' + target.lstrip('/')
        worksheet_path = worksheet_path.replace('xl/xl/', 'xl/')
        sheet_root = ET.fromstring(archive.read(worksheet_path))

        rows = []
        for row in sheet_root.findall(f'.//{{{spreadsheet_ns}}}row'):
            values = {}
            for cell in row.findall(f'{{{spreadsheet_ns}}}c'):
                column = excel_column_index(cell.attrib.get('r', ''))
                cell_type = cell.attrib.get('t')
                value_node = cell.find(f'{{{spreadsheet_ns}}}v')
                if cell_type == 'inlineStr':
                    inline = cell.find(f'{{{spreadsheet_ns}}}is')
                    value = ''.join(node.text or '' for node in inline.iter(f'{{{spreadsheet_ns}}}t')) if inline is not None else ''
                else:
                    value = value_node.text if value_node is not None else ''
                    if cell_type == 's' and value:
                        value = shared_strings[int(value)]
                values[column] = value
            if values:
                rows.append([values.get(i, '') for i in range(max(values) + 1)])

    if not rows:
        return []
    headers = [str(header).strip() for header in rows[0]]
    header_positions = {header.casefold(): index for index, header in enumerate(headers)}

    def get_value(row, header):
        index = header_positions.get(header.casefold())
        return str(row[index]).strip() if index is not None and index < len(row) else ''

    items = []
    for row in rows[1:]:
        container_number = get_value(row, 'Container nummer').upper()
        if not container_number:
            continue
        items.append({
            'container_number': container_number,
            'order_temp': get_value(row, 'Order Temp'),
            'position': get_value(row, 'Positie'),
            'alarms': get_value(row, 'Alarm(en)'),
            'etd': excel_date_display(get_value(row, 'Vertrek')),
            'vessel': get_value(row, 'Vessel'),
            'requested_repair': get_value(row, 'Gevraagd repair'),
            'remarks': get_value(row, 'Opmerkingen')
        })
    return items

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in {'png', 'jpg', 'jpeg', 'gif'}

def create_email_body(report, jobs, alarms, afmelding="", photo_counts=None, ice_procedures=None):
    """Create the original REMS email layout with the updated Job Tasks section."""
    if not report:
        return "<p>Repair Report submitted</p>"

    def safe(value, fallback="N/A"):
        if value is None or value == "":
            return fallback
        return html.escape(str(value))

    def qty(value):
        if value is None or value == "":
            return "N/A"
        try:
            number = float(value)
            return str(int(number)) if number.is_integer() else str(number)
        except (TypeError, ValueError):
            return safe(value)

    afmelding_value = (afmelding or "").strip()
    if afmelding_value.lower() == "nee":
        afmelding_display = '<span style="color:#d60000;font-weight:bold;font-size:16px;">NEE</span>'
    elif afmelding_value.lower() == "ja":
        afmelding_display = '<span style="color:#168a2e;font-weight:bold;">JA</span>'
    else:
        afmelding_display = "N/A"

    html_body = f"""
    <html>
    <head>
        <style>
            body {{ font-family: Arial, sans-serif; line-height: 1.6; }}
            table {{ border-collapse: collapse; width: 100%; margin-bottom: 20px; }}
            th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
            th {{ background-color: #f2f2f2; }}
            .section {{ margin-bottom: 20px; }}
            .section-title {{ font-weight: bold; font-size: 18px; margin-bottom: 10px; }}
        </style>
    </head>
    <body>
        <h2>Repair Report for Container: {safe(report.container_number)}</h2>

        <div class="section">
            <div class="section-title">General Information</div>
            <table>
                <tr><th>Container Number</th><td>{safe(report.container_number)}</td></tr>
                <tr><th>Date</th><td>{safe(report.report_date)}</td></tr>
                <tr><th>Technician</th><td>{safe(report.technician_name)}</td></tr>
                <tr><th>Model</th><td>{safe(report.model)}</td></tr>
                <tr><th>Serial Number</th><td>{safe(report.serial_number)}</td></tr>
                <tr><th>Warranty ID</th><td>{safe(report.warranty_id)}</td></tr>
                <tr><th>Warranty Status</th><td>{safe(report.warranty_status)}</td></tr>
            </table>
        </div>

        <div class="section">
            <div class="section-title">Settings and Readings</div>
            <table>
                <tr><th>Setpoint</th><td>{safe(report.setpoint)} °C</td></tr>
                <tr><th>Vents</th><td>{safe(report.vents)}</td></tr>
                <tr><th>Humidity</th><td>{safe(report.humidity)}</td></tr>
                <tr><th>Ambient</th><td>{safe(report.ambient_temp)} °C</td></tr>
                <tr><th>Supply Temp Before</th><td>{safe(report.supply_temp_before)} °C</td></tr>
                <tr><th>Supply Temp After</th><td>{safe(report.supply_temp_after)} °C</td></tr>
                <tr><th>Return Temp Before</th><td>{safe(report.return_temp_before)} °C</td></tr>
                <tr><th>Return Temp After</th><td>{safe(report.return_temp_after)} °C</td></tr>
                <tr><th>Temperature In Range</th><td>{safe(report.temp_in_range)}</td></tr>
                <tr><th>Afmelding</th><td>{afmelding_display}</td></tr>
            </table>
        </div>

        <div class="section">
            <div class="section-title">Problem Description</div>
            <p>{safe(report.problem_description)}</p>
        </div>

        <div class="section">
            <div class="section-title">Comments</div>
            <p>{safe(report.comments)}</p>
        </div>
    """

    # Updated Job Tasks section only.
    if jobs:
        html_body += """
        <div class="section">
            <div class="section-title">Job Tasks</div>
            <table>
                <tr>
                    <th>Job Code</th>
                    <th>Description</th>
                    <th>Used Part / Bottle</th>
                    <th>Quantity</th>
                    <th>Old Serial</th>
                    <th>New Serial</th>
                    <th>Labor Hours</th>
                </tr>
        """

        total_labor = 0.0

        for job in jobs:
            code = (job.job_code or "").strip().upper()

            try:
                total_labor += float(job.labor_hours or 0)
            except (TypeError, ValueError):
                pass

            if code == "E001II":
                used_part = f"Bottle nr: {safe(job.bottle_number)}"
                old_serial = "—"
                new_serial = "—"
            else:
                part_number = safe(job.part_number, "")
                part_description = safe(job.part_description, "")
                if part_number and part_description:
                    used_part = f"{part_number}<br>{part_description}"
                else:
                    used_part = part_number or part_description or "N/A"
                old_serial = safe(job.old_serial, "—")
                new_serial = safe(job.new_serial, "—")

            html_body += f"""
                <tr>
                    <td>{safe(job.job_code)}</td>
                    <td>{safe(job.description)}</td>
                    <td>{used_part}</td>
                    <td>{qty(job.quantity)}</td>
                    <td>{old_serial}</td>
                    <td>{new_serial}</td>
                    <td>{qty(job.labor_hours)}</td>
                </tr>
            """

        html_body += f"""
                <tr>
                    <td colspan="6" style="text-align:right;font-weight:bold;">Total Labor</td>
                    <td style="font-weight:bold;">{qty(round(total_labor, 2))}</td>
                </tr>
            </table>
        </div>
        """

    if ice_procedures:
        for ice in ice_procedures:
            alarm_row = ""
            if str(ice.get('alarm_active') or '').strip().casefold() == 'yes':
                alarm_row = f"<tr><th>Alarm</th><td>{safe(ice.get('alarm'))}</td></tr>"

            html_body += f"""
        <div class="section">
            <div class="section-title">Ice Procedure</div>
            <table>
                <tr><th>CD26</th><td>{safe(ice.get('cd26'))}</td></tr>
                <tr><th>CD27</th><td>{safe(ice.get('cd27'))}</td></tr>
                <tr><th>Drain hose</th><td>{safe(ice.get('drain_hose'))}</td></tr>
                <tr><th>Upper supply sensor</th><td>{safe(ice.get('upper_supply_sensor'))}</td></tr>
                <tr><th>Lower supply sensor</th><td>{safe(ice.get('lower_supply_sensor'))}</td></tr>
                <tr><th>Alarm Active</th><td>{safe(ice.get('alarm_active'))}</td></tr>
                {alarm_row}
            </table>
        </div>
            """

    if alarms:
        html_body += """
        <div class="section">
            <div class="section-title">Alarms</div>
            <ul>
        """

        for alarm in alarms:
            html_body += f"<li>{safe(alarm.alarm_code)}</li>"

        html_body += "</ul></div>"

    html_body += """
        <div class="section">
            <p>This report was automatically generated by the REMS system.</p>
        </div>
    </body>
    </html>
    """

    return html_body


def send_email(subject, body, attachments, report=None, jobs=None, alarms=None, afmelding="", photo_counts=None, ice_procedures=None):
    SMTP_SERVER = 'smtp.gmail.com'
    SMTP_PORT = 587
    SMTP_USERNAME = os.environ.get('EMAIL_USER')
    SMTP_PASSWORD = os.environ.get('EMAIL_PASS')
    EMAIL_FROM = os.environ.get('EMAIL_FROM', SMTP_USERNAME)
    EMAIL_TO = [address.strip() for address in os.environ.get("EMAIL_TO", "").split(",") if address.strip()]

    msg = MIMEMultipart()
    msg['From'] = EMAIL_FROM
    msg['To'] = ', '.join(EMAIL_TO)
    msg['Subject'] = f"Herstelmelding {subject} - {datetime.now().strftime('%d-%m-%Y')}"

    html_content = create_email_body(report, jobs, alarms, afmelding, photo_counts, ice_procedures)
    msg.attach(MIMEText(html_content, 'html'))

    for filepath in attachments:
        try:
            with open(filepath, 'rb') as f:
                if filepath.lower().endswith(('.png', '.jpg', '.jpeg', '.gif')):
                    img = MIMEImage(f.read())
                    img.add_header('Content-Disposition', 'attachment', filename=os.path.basename(filepath))
                    msg.attach(img)
                else:
                    part = MIMEApplication(f.read(), Name=os.path.basename(filepath))
                    part['Content-Disposition'] = f'attachment; filename="{os.path.basename(filepath)}"'
                    msg.attach(part)
        except Exception as e:
            app.logger.error(f"Failed to attach {filepath}: {str(e)}")

    with smtplib.SMTP(SMTP_SERVER, SMTP_PORT) as smtp:
        smtp.starttls()
        smtp.login(SMTP_USERNAME, SMTP_PASSWORD)
        smtp.sendmail(EMAIL_FROM, EMAIL_TO, msg.as_string())
        app.logger.info(f"Email sent to: {EMAIL_TO}")



def _weekly_freon_period(reference_date=None):
    """Return the previous Monday-Sunday period."""
    today = reference_date or datetime.utcnow().date()
    current_monday = today - timedelta(days=today.weekday())
    start_date = current_monday - timedelta(days=7)
    end_date = current_monday - timedelta(days=1)
    return start_date, end_date


def _build_freon_workbook(rows, start_date, end_date):
    """Create the weekly Freon registration workbook in memory."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Freon Registration"

    sheet.append(["Date", "Container nr", "Bottle nr", "Quantity", "Technician"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for job, report in rows:
        sheet.append([
            report.report_date,
            report.container_number or "",
            job.bottle_number or "",
            float(job.quantity or 0),
            report.technician_name or ""
        ])

    for cell in sheet["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for cell in sheet["D"][1:]:
        cell.number_format = "0.00"

    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 12
    sheet.column_dimensions["E"].width = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"Freon_Registration_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    return filename, output.getvalue()


def send_weekly_freon_registration(reference_date=None):
    """Email unsent E001II registrations from the previous Monday-Sunday."""
    start_date, end_date = _weekly_freon_period(reference_date)

    rows = (
        db.session.query(RepairJob, RepairReport)
        .join(RepairReport, RepairJob.report_id == RepairReport.id)
        .filter(
            RepairJob.job_code == "E001II",
            RepairReport.report_date >= start_date,
            RepairReport.report_date <= end_date,
            RepairJob.freon_sent_at.is_(None)
        )
        .order_by(RepairReport.report_date.asc(), RepairReport.id.asc(), RepairJob.id.asc())
        .all()
    )

    smtp_username = os.environ.get("EMAIL_USER")
    smtp_password = os.environ.get("EMAIL_PASS")
    email_from = os.environ.get("EMAIL_FROM", smtp_username)
    recipient_setting = os.environ.get("FREON_EMAIL_TO") or os.environ.get("EMAIL_TO", "")
    email_to = [address.strip() for address in recipient_setting.split(",") if address.strip()]

    if not smtp_username or not smtp_password:
        raise RuntimeError("Gmail credentials are not configured")
    if not email_to:
        raise RuntimeError("FREON_EMAIL_TO is not configured")

    def esc(value):
        return html.escape(str(value if value is not None else ""))

    table_rows = ""
    for job, report in rows:
        quantity = float(job.quantity or 0)
        quantity_text = str(int(quantity)) if quantity.is_integer() else f"{quantity:g}"
        table_rows += f"""
        <tr>
          <td style="padding:7px 10px;border:1px solid #d9d9d9;">{report.report_date.strftime('%d/%m/%Y')}</td>
          <td style="padding:7px 10px;border:1px solid #d9d9d9;">{esc(report.container_number)}</td>
          <td style="padding:7px 10px;border:1px solid #d9d9d9;">{esc(job.bottle_number)}</td>
          <td style="padding:7px 10px;border:1px solid #d9d9d9;text-align:right;">{esc(quantity_text)}</td>
          <td style="padding:7px 10px;border:1px solid #d9d9d9;">{esc(report.technician_name)}</td>
        </tr>"""

    if not table_rows:
        table_rows = '<tr><td colspan="5" style="padding:10px;border:1px solid #d9d9d9;">No R134A registrations for this period.</td></tr>'

    body = f"""
    <html><body style="font-family:Arial,sans-serif;color:#222;">
      <p>Hi,</p>
      <p>Please find below the weekly Freon registration for <strong>{start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')}</strong>.</p>
      <table style="border-collapse:collapse;font-size:13px;">
        <thead><tr style="background:#f2f2f2;">
          <th style="padding:7px 10px;border:1px solid #d9d9d9;">Date</th>
          <th style="padding:7px 10px;border:1px solid #d9d9d9;">Container nr</th>
          <th style="padding:7px 10px;border:1px solid #d9d9d9;">Bottle nr</th>
          <th style="padding:7px 10px;border:1px solid #d9d9d9;">Quantity</th>
          <th style="padding:7px 10px;border:1px solid #d9d9d9;">Technician</th>
        </tr></thead>
        <tbody>{table_rows}</tbody>
      </table>
      <p>Kind regards,<br>REMS</p>
    </body></html>
    """

    msg = MIMEMultipart()
    msg["From"] = email_from
    msg["To"] = ", ".join(email_to)
    msg["Subject"] = f"REMS - Weekly Freon Registration - {start_date.strftime('%d-%m-%Y')} to {end_date.strftime('%d-%m-%Y')}"
    msg.attach(MIMEText(body, "html"))

    if rows:
        filename, workbook_bytes = _build_freon_workbook(rows, start_date, end_date)
        attachment = MIMEApplication(
            workbook_bytes,
            _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        attachment.add_header("Content-Disposition", "attachment", filename=filename)
        msg.attach(attachment)

    with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
        smtp.starttls()
        smtp.login(smtp_username, smtp_password)
        smtp.sendmail(email_from, email_to, msg.as_string())

    # Only mark records after SMTP confirms the message was sent.
    sent_at = datetime.utcnow()
    for job, _report in rows:
        job.freon_sent_at = sent_at
    db.session.commit()

    app.logger.info(
        "Weekly Freon registration sent: %s rows, %s to %s",
        len(rows), start_date, end_date
    )
    return {
        "status": "success",
        "count": len(rows),
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat()
    }


@app.route('/api/weekly-freon-registration', methods=['POST'])
def weekly_freon_registration_route():
    """Protected endpoint intended for the Render Cron Job."""
    expected_secret = os.environ.get("CRON_SECRET", "").strip()
    supplied_secret = request.headers.get("X-Cron-Secret", "").strip()
    if not expected_secret or supplied_secret != expected_secret:
        return jsonify({"status": "error", "message": "Unauthorized"}), 401

    try:
        result = send_weekly_freon_registration()
        return jsonify(result), 200
    except Exception as exc:
        db.session.rollback()
        app.logger.error("Weekly Freon registration failed: %s", exc, exc_info=True)
        return jsonify({"status": "error", "message": str(exc)}), 500

def send_repair_zone_email(container_number, technician_name, current_position):
    smtp_server = 'smtp.gmail.com'
    smtp_port = 587
    smtp_username = os.environ.get('EMAIL_USER')
    smtp_password = os.environ.get('EMAIL_PASS')
    email_from = os.environ.get('EMAIL_FROM', smtp_username)
    recipient_setting = os.environ.get('REPAIR_ZONE_EMAIL_TO') or os.environ.get('EMAIL_TO', '')
    email_to = [address.strip() for address in recipient_setting.split(',') if address.strip()]

    if not smtp_username or not smtp_password:
        raise RuntimeError('Gmail credentials are not configured')
    if not email_to:
        raise RuntimeError('REPAIR_ZONE_EMAIL_TO is not configured')

    safe_container = html.escape(container_number)
    safe_technician = html.escape(technician_name)
    safe_position = html.escape(current_position)
    requested_at = datetime.utcnow().strftime('%d-%m-%Y %H:%M UTC')

    message = MIMEMultipart('alternative')
    message['From'] = email_from
    message['To'] = ', '.join(email_to)
    message['Subject'] = f"REMS Repair Zone Request - {container_number}"
    message.attach(MIMEText(f"""
    <html><body style="font-family:Arial,sans-serif;color:#222;">
      <h2 style="color:#003366;">Reefer Repair Zone Request</h2>
      <table style="border-collapse:collapse;">
        <tr><td style="padding:6px 18px 6px 0;"><strong>Container</strong></td><td>{safe_container}</td></tr>
        <tr><td style="padding:6px 18px 6px 0;"><strong>Current position</strong></td><td>{safe_position}</td></tr>
        <tr><td style="padding:6px 18px 6px 0;"><strong>Requested by</strong></td><td>{safe_technician}</td></tr>
        <tr><td style="padding:6px 18px 6px 0;"><strong>Date/time</strong></td><td>{requested_at}</td></tr>
      </table>
      <p>Please bring this reefer to the repair zone.</p>
      <p style="color:#666;font-size:12px;">This request was automatically generated by REMS.</p>
    </body></html>
    """, 'html'))

    with smtplib.SMTP(smtp_server, smtp_port) as smtp:
        smtp.starttls()
        smtp.login(smtp_username, smtp_password)
        smtp.sendmail(email_from, email_to, message.as_string())
        app.logger.info("Repair-zone request sent for %s to %s", container_number, email_to)

# ===================================
# Run App
# ===================================
if __name__ == '__main__':
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
